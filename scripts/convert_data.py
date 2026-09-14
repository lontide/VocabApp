import openpyxl
import json
import os

excel_path = os.path.join(os.path.dirname(__file__), '..', 'wordlist.xlsx')
output_json_path = os.path.join(os.path.dirname(__file__), '..', 'src', 'data', 'wordlist.json')

wb = openpyxl.load_workbook(excel_path, data_only=True)

# Process all '총정리_X탄' sheets
sheets = [s for s in wb.sheetnames if s.startswith('총정리_')]
# Sort sheets by number
sheets.sort(key=lambda x: int(x.replace('총정리_', '').replace('탄', '')))

all_items = []
rounds_info = []

for s in sheets:
    round_num = int(s.replace('총정리_', '').replace('탄', ''))
    sheet = wb[s]
    rows = list(sheet.iter_rows(values_only=True))
    sheet_items = []
    
    for r in rows[1:]:
        if not r or not any(r):
            continue
        no = r[0]
        cat = str(r[1] or '').strip()
        word = str(r[2] or '').strip()
        meaning = str(r[3] or '').strip()
        ex_en = str(r[4] or '').strip()
        ex_ko = str(r[5] or '').strip()
        link = str(r[6] or '').strip() if len(r) > 6 and r[6] else ''
        
        item = {
            'id': int(no) if isinstance(no, (int, float)) else len(all_items) + 1,
            'round': round_num,
            'category': cat,
            'word': word,
            'meaning': meaning,
            'ex_en': ex_en,
            'ex_ko': ex_ko,
            'link': link
        }
        sheet_items.append(item)
        all_items.append(item)
        
    categories = sorted(list(set(item['category'] for item in sheet_items if item['category'])))
    rounds_info.append({
        'round': round_num,
        'title': f'총정리 {round_num}탄',
        'count': len(sheet_items),
        'categories': categories,
        'itemIds': [item['id'] for item in sheet_items]
    })

data = {
    'totalCount': len(all_items),
    'rounds': rounds_info,
    'items': all_items
}

with open(output_json_path, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print(f"Successfully converted {len(all_items)} expressions across {len(rounds_info)} rounds to {output_json_path}")
